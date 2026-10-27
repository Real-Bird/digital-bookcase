import copy
import ctypes
import json
import os
import sys
import uuid
from pathlib import Path
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from domain import normalize_isbn, parse_results
from storage import Store, StoreError, COLUMNS
from settings import Settings
import app as module


@pytest.fixture
def store(tmp_path):
    return Store(tmp_path / '소장도서.xlsx', retention=2)


def book(title='아몬드', isbn='9788936434267'):
    return module.Book(title=title, isbn=isbn).model_dump()


def save(store, data=None, duplicate='cancel', op=None, revision=None):
    return store.mutate('save', op or str(uuid.uuid4()), revision or store.revision(), {'book': data or book(), 'duplicate':duplicate})


def test_isbn_conversion_and_rejection():
    assert normalize_isbn('89-364-3426-8') == '9788936434267'
    assert normalize_isbn('0-8044-2957-X') == '9780804429573'
    for value in ['9788936434268','9790123456785','123','9780000000003','893643426X']:
        with pytest.raises(ValueError): normalize_isbn(value)
    assert normalize_isbn('', allow_empty=True) == ''


def test_live_response_shapes():
    fixture = json.loads((Path(__file__).parent / 'api-fixtures.json').read_text('utf-8-sig'))
    one = parse_results(fixture['normal'], '9788936434267')
    assert one['items'][0]['title'].startswith('아몬드')
    assert one['items'][0]['detail_link'].startswith('https://www.nl.go.kr/NL/')
    assert '#viewKey=' in one['items'][0]['detail_link']
    assert parse_results(fixture['no_results'],'9798765432105')['items'] == []
    mismatch = parse_results(fixture['isbn_mismatch'],'9780000000002')
    assert mismatch['rejected'] == 4
    assert mismatch['items'] == []
    assert len(parse_results(fixture['multiple_results'],'9781234567897')['items']) == 3
    with pytest.raises(ValueError, match='010'): parse_results(fixture['missing_key'],'9788936434267')


def test_text_and_metadata_persist(store):
    data=book('=HYPERLINK("https://example.com")')
    data['memo']='한글 메모\n둘째 줄'
    result=save(store,data)
    loaded=Store(store.path).snapshot()['books'][0]
    assert loaded['isbn'] == data['isbn'] and loaded['memo'] == data['memo']
    wb=load_workbook(store.path)
    assert wb['소장도서']['E2'].data_type == 's'
    assert wb['소장도서']['B2'].number_format == '@'
    assert wb['_operations'].sheet_state == 'hidden'
    assert result['book']['id'] == loaded['id']


def test_duplicate_increment_idempotent_after_restart(store):
    save(store)
    with pytest.raises(StoreError) as exc: save(store)
    assert exc.value.details['duplicate']['quantity'] == 1
    op=str(uuid.uuid4());rev=store.revision();payload={'book':book(), 'duplicate':'increment'}
    first=store.mutate('save',op,rev,payload)
    replay=Store(store.path).mutate('save',op,rev,payload)
    assert first['book']['quantity'] == replay['book']['quantity'] == 2
    assert replay['replayed']
    assert Store(store.path).snapshot()['books'][0]['quantity'] == 2
    payload['book']['quantity']=3
    with pytest.raises(StoreError): store.mutate('save',op,rev,payload)


def test_external_change_and_failed_replace_keep_original(store):
    save(store);old=store.revision()
    wb=load_workbook(store.path);wb['소장도서']['E2']='외부 수정';wb.save(store.path)
    with pytest.raises(StoreError,match='목록이 변경'): save(store,revision=old)
    before=store.path.read_bytes()
    with patch('storage.os.replace',side_effect=PermissionError):
        with pytest.raises(StoreError) as exc: save(store,book('다른 책',''))
    assert exc.value.code == 423 and store.path.read_bytes() == before
    assert not list(store.path.parent.glob('.saving-*'))


def test_external_edit_during_commit_detected(store):
    save(store);before=store.path.read_bytes()
    original=store.revision
    count=[0]
    def fake_revision():
        count[0]+=1
        return original() if count[0]<3 else 'changed'
    with patch.object(store,'revision',side_effect=fake_revision):
        with pytest.raises(StoreError): save(store,book('새 책',''))
    assert store.path.read_bytes() == before


def test_delete_and_restore_receipts_and_retention(store):
    first=save(store)['book']
    save(store,book('다음 책',''))
    backup=store.backups()[0]['name']
    restore_op=str(uuid.uuid4());rev=store.revision()
    store.mutate('restore',restore_op,rev,{'name':backup})
    assert len(store.snapshot()['books']) == 1
    store.mutate('restore',restore_op,rev,{'name':backup})
    delete_op=str(uuid.uuid4());rev=store.revision()
    store.mutate('delete',delete_op,rev,{'id':first['id']})
    assert not store.snapshot()['books']
    assert store.mutate('delete',delete_op,rev,{'id':first['id']})['replayed']
    assert len(store.backups()) == 2


def test_refuse_foreign_workbook_and_bad_quantity(store):
    from openpyxl import Workbook
    wb=Workbook();wb.save(store.path)
    with pytest.raises(StoreError,match='전용'): store.snapshot()
    wb=store.empty();ws=wb['소장도서'];ws.append(['abc','','','','제목','','','','잘못된 권수']);wb.save(store.path)
    with pytest.raises(StoreError,match='권수'): store.snapshot()


@pytest.mark.skipif(os.name != 'nt', reason='Windows file sharing')
def test_windows_excel_style_lock_preserves_file(store):
    from ctypes import wintypes
    save(store)
    original=store.path.read_bytes()
    kernel=ctypes.WinDLL('kernel32', use_last_error=True)
    kernel.CreateFileW.argtypes=[wintypes.LPCWSTR,wintypes.DWORD,wintypes.DWORD,ctypes.c_void_p,wintypes.DWORD,wintypes.DWORD,wintypes.HANDLE]
    kernel.CreateFileW.restype=wintypes.HANDLE
    kernel.CloseHandle.argtypes=[wintypes.HANDLE]
    handle=kernel.CreateFileW(str(store.path),0x80000000,3,None,3,0,None)
    assert handle != wintypes.HANDLE(-1).value
    try:
        with pytest.raises(StoreError) as exc: save(store,book('잠금 중 새 책',''))
        assert exc.value.code==423
        assert store.path.read_bytes()==original
    finally:
        kernel.CloseHandle(handle)


def test_concurrent_retry_receipt_is_once(store):
    from concurrent.futures import ThreadPoolExecutor
    op=str(uuid.uuid4());rev=store.revision();payload={'book':book(),'duplicate':'cancel'}
    with ThreadPoolExecutor(max_workers=2) as pool:
        results=list(pool.map(lambda _:store.mutate('save',op,rev,payload),range(2)))
    assert sum(bool(r['replayed']) for r in results)==1
    assert len(store.snapshot()['books'])==1


@pytest.mark.skipif(os.name != 'nt', reason='Windows DPAPI')
def test_key_is_encrypted_and_persistent(tmp_path):
    settings=Settings(tmp_path)
    settings.save(key='test-private-key')
    assert 'test-private-key' not in settings.file.read_text('utf-8')
    assert Settings(tmp_path).key() == 'test-private-key'
    assert settings.public()['key_configured']
    assert 'key_encrypted' not in settings.public()


@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setattr(module,'settings',Settings(tmp_path))
    monkeypatch.setattr(module,'store',Store(tmp_path / 'books.xlsx'))
    monkeypatch.setattr(module,'cache',{})
    with TestClient(module.app) as client: yield client


def test_http_guards_and_validation(client):
    assert client.get('/api/bootstrap').status_code == 200
    assert client.post('/api/books',json={}).status_code == 403
    headers={'x-session-token':module.session_token}
    assert client.post('/api/books',headers=headers,json={}).status_code == 422
    assert client.get('/api/bootstrap',headers={'host':'evil.example'}).status_code == 403
    assert client.post('/api/open-folder',headers={**headers,'origin':'https://evil.example'},json={}).status_code == 403
    response=client.get('/')
    assert response.status_code == 200 and 'frame-ancestors' in response.headers['content-security-policy']
    response=client.post('/api/lookup',headers=headers,json={'isbn':'9788936434268'})
    assert response.status_code == 400


def test_end_to_end_http_save_edit_delete(client):
    headers={'x-session-token':module.session_token}
    data=client.get('/api/state').json()
    request={'book':book(),'op_id':str(uuid.uuid4()),'revision':data['revision']}
    response=client.post('/api/books',headers=headers,json=request)
    assert response.status_code == 200
    result=response.json()
    assert client.post('/api/books',headers=headers,json=request).json()['replayed']
    result['book']['memo']='수정 메모'
    edited=client.post('/api/books',headers=headers,json={'book':result['book'],'op_id':str(uuid.uuid4()),'revision':result['revision']}).json()
    assert edited['book']['memo']=='수정 메모'
    deleted=client.post('/api/delete',headers=headers,json={'id':edited['book']['id'],'op_id':str(uuid.uuid4()),'revision':edited['revision']})
    assert deleted.status_code==200
    assert client.get('/api/books').json()['books']==[]
