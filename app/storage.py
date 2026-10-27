"""Excel is the authoritative store; operation receipts commit in the same file."""
import hashlib
import json
import os
import shutil
import tempfile
import threading
import uuid
from datetime import datetime, timezone, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

COLUMNS = [('id','등록ID'), ('isbn','ISBN13'), ('raw_input','원본입력'), ('response_isbn','응답ISBN'),
           ('title','제목'), ('author','저자'), ('publisher','출판사'), ('year','발행년도'),
           ('quantity','권수'), ('location','소장위치'), ('memo','메모'), ('source','정보출처'),
           ('lookup_status','조회상태'), ('control_no','NL제어번호'), ('nl_id','NL종키'),
           ('call_no','NL청구기호'), ('kdc_code','KDC코드'), ('kdc_name','KDC명칭'),
           ('detail_link','상세링크'), ('created_at','등록일시'), ('updated_at','수정일시'), ('looked_up_at','조회일시')]
KST = timezone(timedelta(hours=9))


def now():
    return datetime.now(KST).isoformat(timespec='seconds')


class StoreError(Exception):
    def __init__(self, message, code=409, details=None):
        super().__init__(message)
        self.code = code
        self.details = details


class Store:
    def __init__(self, path, retention=20):
        self.path = Path(path)
        self.retention = retention
        self.lock = threading.RLock()

    def revision(self):
        return hashlib.sha256(self.path.read_bytes()).hexdigest() if self.path.exists() else 'new'

    def empty(self):
        wb = Workbook()
        ws = wb.active
        ws.title = '소장도서'
        ws.append([label for _, label in COLUMNS])
        meta = wb.create_sheet('_meta')
        meta.append(['book-scanner', 1])
        meta.sheet_state = 'hidden'
        ops = wb.create_sheet('_operations')
        ops.append(['작업ID', '요청해시', '결과JSON'])
        ops.sheet_state = 'hidden'
        return wb

    def validate(self, wb):
        if not {'소장도서','_meta','_operations'}.issubset(wb.sheetnames) or wb['_meta'].cell(1,1).value != 'book-scanner' or wb['_meta'].cell(1,2).value != 1:
            raise StoreError('앱 전용 Excel 형식이 아닙니다. 새 파일명을 지정하세요.', 400)
        if [c.value for c in wb['소장도서'][1]] != [label for _,label in COLUMNS]:
            raise StoreError('Excel 열 구성이 변경되었습니다. 원래 열 구성을 복구하세요.', 400)
        ids = set()
        for row in wb['소장도서'].iter_rows(min_row=2, values_only=True):
            if not any(v is not None for v in row):
                continue
            record = dict(zip((k for k,_ in COLUMNS), row))
            if not record['id'] or record['id'] in ids or not isinstance(record['quantity'], int) or isinstance(record['quantity'], bool) or record['quantity'] < 1:
                raise StoreError('Excel 등록ID 또는 권수 값이 올바르지 않습니다.', 400)
            ids.add(record['id'])

    def read(self, path=None):
        target = Path(path or self.path)
        if not target.exists():
            return self.empty()
        try:
            wb = load_workbook(target)
            self.validate(wb)
            return wb
        except StoreError:
            raise
        except PermissionError:
            raise StoreError('Excel 파일에 접근할 수 없습니다. 파일을 닫고 다시 시도하세요.', 423)
        except Exception:
            raise StoreError('Excel 파일을 읽을 수 없습니다. 파일 상태나 백업을 확인하세요.', 400)

    @staticmethod
    def records(wb):
        return [{k: (value if value is not None else '') for (k,_),value in zip(COLUMNS,row)}
                for row in wb['소장도서'].iter_rows(min_row=2, values_only=True) if row[0]]

    def snapshot(self):
        with self.lock:
            before = self.revision()
            wb = self.read()
            if before != self.revision():
                raise StoreError('다른 프로그램에서 파일을 수정했습니다. 목록을 다시 읽으세요.')
            return {'books': self.records(wb), 'revision': before, 'path': str(self.path)}

    def backup_dir(self):
        return self.path.parent / 'backups' / self.path.stem

    def backups(self):
        folder = self.backup_dir()
        return [{'name': p.name, 'size': p.stat().st_size} for p in sorted(folder.glob('*.xlsx'), reverse=True)] if folder.exists() else []

    def write(self, wb, expected):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        ws = wb['소장도서']
        ws.freeze_panes = 'E2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='254F48')
            cell.alignment = Alignment(vertical='center')
        ws.row_dimensions[1].height = 28
        for i,(key,_) in enumerate(COLUMNS, 1):
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(i)].width = 45 if key in ('title','memo','detail_link') else 23
        fd, temp = tempfile.mkstemp(prefix='.saving-', suffix='.xlsx', dir=self.path.parent)
        os.close(fd)
        try:
            wb.save(temp)
            check = load_workbook(temp)
            self.validate(check)
            check.close()
            with open(temp, 'r+b') as handle:
                os.fsync(handle.fileno())
            if expected != self.revision():
                raise StoreError('Excel 파일이 외부에서 수정되었습니다. 목록을 새로고침하세요.')
            if self.path.exists():
                folder = self.backup_dir()
                folder.mkdir(parents=True, exist_ok=True)
                stamp = datetime.now(KST).strftime('%Y%m%d-%H%M%S-%f')
                shutil.copy2(self.path, folder / (stamp + '.xlsx'))
                if expected != self.revision():
                    raise StoreError('백업 도중 외부 변경이 감지되었습니다. 목록을 다시 읽으세요.')
            os.replace(temp, self.path)
        except PermissionError:
            raise StoreError('Excel 파일이 열려 있거나 쓰기 권한이 없습니다. 파일을 닫고 다시 저장하세요.', 423)
        except StoreError:
            raise
        except OSError:
            raise StoreError('저장하지 못했습니다. 저장 공간과 폴더 권한을 확인하세요.', 500)
        finally:
            if os.path.exists(temp):
                os.unlink(temp)
        # Pruning failure must never turn a committed operation into a failed save.
        for backup in self.backups()[self.retention:]:
            try:
                (self.backup_dir() / backup['name']).unlink()
            except OSError:
                pass

    def mutate(self, action, op_id, expected, payload):
        digest = hashlib.sha256(json.dumps([action,payload], ensure_ascii=False, sort_keys=True).encode()).hexdigest()
        with self.lock:
            wb = self.read()
            ops = wb['_operations']
            for row in ops.iter_rows(min_row=2, values_only=True):
                if row[0] == op_id:
                    if row[1] != digest:
                        raise StoreError('같은 작업ID의 요청 내용이 바뀌었습니다. 새 작업으로 다시 시도하세요.')
                    return {**json.loads(row[2]), 'revision': self.revision(), 'replayed': True}
            if expected != self.revision():
                raise StoreError('목록이 변경되었습니다. 새로고침 후 다시 저장하세요.')
            books = self.records(wb)
            if action == 'save':
                book = dict(payload['book'])
                target = next((x for x in books if x['id'] == book.get('id')), None)
                if book.get('id') and not target:
                    raise StoreError('수정할 도서가 없습니다.', 404)
                duplicate = next((x for x in books if book['isbn'] and x['isbn'] == book['isbn'] and x['id'] != book.get('id')), None)
                if duplicate and payload['duplicate'] == 'cancel':
                    raise StoreError('이미 등록된 ISBN입니다.', 409, {'duplicate': duplicate})
                if duplicate and payload['duplicate'] == 'increment':
                    if book.get('id'):
                        raise StoreError('수정 도서의 ISBN이 다른 등록 도서와 겹칩니다.')
                    duplicate['quantity'] += book['quantity']
                    duplicate['updated_at'] = now()
                    result = duplicate
                elif target:
                    book['created_at'] = target['created_at']
                    book['updated_at'] = now()
                    books[books.index(target)] = book
                    result = book
                else:
                    book.update(id=str(uuid.uuid4()), created_at=now(), updated_at=now())
                    books.append(book)
                    result = book
                response = {'book': result, 'message': '저장 완료'}
            elif action == 'delete':
                target = next((x for x in books if x['id'] == payload['id']), None)
                if not target:
                    raise StoreError('삭제할 도서가 없습니다.', 404)
                books.remove(target)
                response = {'message': '등록 취소 완료'}
            elif action == 'restore':
                name = payload['name']
                if name not in [b['name'] for b in self.backups()]:
                    raise StoreError('백업을 찾을 수 없습니다.', 404)
                restored = self.read(self.backup_dir() / name)
                books = self.records(restored)
                response = {'message': '백업 복원 완료'}
            else:
                raise StoreError('지원하지 않는 작업입니다.', 400)
            ws = wb['소장도서']
            if ws.max_row > 1:
                ws.delete_rows(2, ws.max_row-1)
            for book in books:
                ws.append([book.get(key, '') for key,_ in COLUMNS])
                for cell in ws[ws.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
                        cell.number_format = '@'
            ops.append([op_id, digest, json.dumps(response, ensure_ascii=False)])
            for cell in ops[ops.max_row]:
                cell.data_type = 's'
            self.write(wb, expected)
            return {**response, 'revision': self.revision(), 'replayed': False}
